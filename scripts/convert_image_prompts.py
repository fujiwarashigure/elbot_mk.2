#!/usr/bin/env python3
"""Convert GPT_Image_Prompts_大全.xlsx into the embedded prompt library JSON.

Usage:
    python scripts/convert_image_prompts.py \
        --input "GPT_Image_Prompts_大全.xlsx" \
        --output internal/imagegen/prompts/library.json
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl

SPLIT_RE = re.compile(r"[/、,，;；|｜\s]+")
VAR_RE = re.compile(r"\[[^\]]*\]")
VAR_LIST_RE = re.compile(r"\[([^\]]+)\]")
NEGATIVE_SLUGS = {
    "画质": "quality",
    "曝光与色彩": "exposure",
    "人体结构": "anatomy",
    "面部": "face",
    "文字与标识": "text",
    "构图": "composition",
    "透视与空间": "perspective",
    "重复与克隆": "duplication",
    "风格问题": "style",
    "安全与合规": "safety",
}


def cell(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def split_keys(*values: str, min_len: int = 2) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for value in values:
        for token in SPLIT_RE.split(value):
            token = token.strip(" ．.。:：()（）[]【】\"'")
            if len(token) < min_len or token in seen:
                continue
            seen.add(token)
            out.append(token)
    return out


WORD_RE = re.compile(r"[A-Za-z][A-Za-z0-9'-]+")
WORD_STOPWORDS = {
    "a", "an", "the", "and", "or", "of", "in", "on", "at", "to", "for", "with", "without",
    "by", "from", "as", "is", "are", "be", "into", "over", "under", "up", "down", "out",
    "this", "that", "these", "those", "it", "its", "his", "her", "their", "your", "you",
    "very", "highly", "more", "most", "no", "not", "without", "using", "use", "keep",
    "unchanged", "image", "photo", "picture", "generate", "create",
}


def split_words(text: str, min_len: int = 3) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for token in WORD_RE.findall(text):
        word = token.strip("'-").lower()
        if len(word) < min_len or word in WORD_STOPWORDS or word in seen:
            continue
        seen.add(word)
        out.append(word)
    return out


def split_list(value: str) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for token in value.split(","):
        token = token.strip()
        if not token or token in seen:
            continue
        seen.add(token)
        out.append(token)
    return out


STOPWORDS = {"of", "in", "on", "with", "and", "for", "to", "at", "from", "by", "the", "a", "an", "into"}
CANONICAL_USE_KEYS = [
    "头像", "壁纸", "封面", "海报", "主图", "详情页", "头图", "分镜", "横幅", "插画",
    "图标", "logo", "ui", "贺卡", "请柬", "明信片", "表情包", "三视图", "信息图",
]


def clean_anchor(text: str) -> str:
    words = text.split()
    while words and words[0].lower() in STOPWORDS:
        words.pop(0)
    while words and words[-1].lower() in STOPWORDS:
        words.pop()
    return " ".join(words).strip(" .,;:，。；：、!?！？")


def canonical_keys(text: str) -> list[str]:
    folded = text.lower()
    return [key for key in CANONICAL_USE_KEYS if key.lower() in folded]


def anchors_from_prompt(prompt: str) -> list[str]:
    stripped = VAR_RE.sub(" ", prompt)
    out: list[str] = []
    seen: set[str] = set()
    for part in stripped.split(","):
        part = clean_anchor(" ".join(part.split()))
        if len(part) < 4 or part in seen:
            continue
        if not any(ch.isalpha() for ch in part):
            continue
        if len(part.split()) == 1 and len(part) < 8:
            continue
        seen.add(part)
        out.append(part)
    return out


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", default="GPT_Image_Prompts_大全.xlsx")
    parser.add_argument("--output", default="internal/imagegen/prompts/library.json")
    args = parser.parse_args()

    source = Path(args.input)
    if not source.exists():
        print(f"input not found: {source}", file=sys.stderr)
        return 1
    book = openpyxl.load_workbook(source, data_only=True, read_only=True)

    library = {
        "version": 2,
        "source": source.name,
        "templates": [],
        "entries": [],
        "keywords": {},
        "keyword_terms": {},
        "negatives": [],
        "presets": [],
        "tips": [],
        "vocabulary": [],
    }

    if "万能模板" in book.sheetnames:
        for row in list(book["万能模板"].iter_rows(values_only=True))[1:]:
            if not cell(row[4]):
                continue
            library["templates"].append({
                "type": cell(row[1]),
                "name": cell(row[2]),
                "scene": cell(row[3]),
                "template": cell(row[4]),
                "variables": VAR_LIST_RE.findall(cell(row[5])),
            })

    if "全部提示词汇总" in book.sheetnames:
        for row in list(book["全部提示词汇总"].iter_rows(values_only=True))[1:]:
            prompt = cell(row[4])
            if not prompt:
                continue
            category = cell(row[1])
            scene = cell(row[2])
            description = cell(row[3])
            variables = VAR_LIST_RE.findall(cell(row[5]))
            library["entries"].append({
                "category": category,
                "scene": scene,
                "description": description,
                "prompt": prompt,
                "variables": variables,
                "anchors": anchors_from_prompt(prompt),
                "terms": split_words(VAR_RE.sub(" ", prompt)),
                "keys": split_keys(category, scene),
                "generic_keys": canonical_keys(category + scene),
            })

    if "关键词词库" in book.sheetnames:
        for row in list(book["关键词词库"].iter_rows(values_only=True))[1:]:
            name = cell(row[0])
            if not name:
                continue
            items = split_list(cell(row[1]))
            library["keywords"][name] = items
            library["keyword_terms"][name] = split_words(" ".join(items))

    if "负面提示词" in book.sheetnames:
        for index, row in enumerate(list(book["负面提示词"].iter_rows(values_only=True))[1:]):
            name = cell(row[1])
            if not name:
                continue
            items = split_list(cell(row[2]))
            library["negatives"].append({
                "key": NEGATIVE_SLUGS.get(name, f"n{index}"),
                "category": name,
                "items": items,
                "terms": split_words(" ".join(items)),
            })

    if "参数速查" in book.sheetnames:
        for row in list(book["参数速查"].iter_rows(values_only=True))[1:]:
            use = cell(row[1])
            if not use:
                continue
            scene = cell(row[3])
            library["presets"].append({
                "use": use,
                "ratio": cell(row[2]),
                "scene": scene,
                "supplement": cell(row[4]),
                "keys": split_keys(use) + canonical_keys(use),
                "scene_keys": split_keys(scene) + canonical_keys(scene),
            })

    if "进阶技巧" in book.sheetnames:
        for row in list(book["进阶技巧"].iter_rows(values_only=True))[1:]:
            title = cell(row[2])
            if not title:
                continue
            library["tips"].append({
                "module": cell(row[1]),
                "title": title,
                "detail": cell(row[3]),
                "example": cell(row[4]),
            })

    counts: dict[str, int] = {}
    for entry in library["entries"]:
        for word in entry["terms"]:
            counts[word] = counts.get(word, 0) + 1
    for items in library["keyword_terms"].values():
        for word in items:
            counts[word] = counts.get(word, 0) + 1
    for group in library["negatives"]:
        for word in group["terms"]:
            counts[word] = counts.get(word, 0) + 1
    library["vocabulary"] = [
        {"word": word, "count": count}
        for word, count in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))[:2000]
    ]

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(library, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(
        f"wrote {output} entries={len(library['entries'])} templates={len(library['templates'])} "
        f"keywords={len(library['keywords'])} negatives={len(library['negatives'])} "
        f"presets={len(library['presets'])} tips={len(library['tips'])} bytes={output.stat().st_size}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
